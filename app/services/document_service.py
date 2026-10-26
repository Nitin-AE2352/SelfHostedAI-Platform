
from pathlib import Path
import hashlib
import io
import re

import chromadb
import requests
from docx import Document
from pypdf import PdfReader
from pptx import Presentation

from app.core.config import settings


# Project folders
PROJECT_ROOT = Path(__file__).resolve().parents[2]
DATA_DIR = PROJECT_ROOT / "data"
DB_DIR = DATA_DIR / "chroma_db"

DATA_DIR.mkdir(parents=True, exist_ok=True)
DB_DIR.mkdir(parents=True, exist_ok=True)

# Persistent local vector database
chroma_client = chromadb.PersistentClient(path=str(DB_DIR))

collection = chroma_client.get_or_create_collection(
    name="platform_documents",
    metadata={"hnsw:space": "cosine"},
)

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".pptx", ".txt"}
MAX_FILE_BYTES = settings.max_upload_mb * 1024 * 1024

CHUNK_SIZE = 900
CHUNK_OVERLAP = 150
EMBED_MODEL = "nomic-embed-text"


def clean_text(text: str) -> str:
    """Normalize extracted document text."""
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_pages(
    filename: str,
    content: bytes,
) -> list[dict]:
    """
    Extract text and source locations.

    PDF: preserve the original page number.
    PPTX: preserve the slide number.
    DOCX/TXT: use document-level locations.
    """
    extension = Path(filename).suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {extension}")

    if not content:
        raise ValueError("The uploaded file is empty.")

    if len(content) > MAX_FILE_BYTES:
        raise ValueError(
            f"File exceeds the {settings.max_upload_mb} MB limit."
        )

    sections = []

    if extension == ".pdf":
        reader = PdfReader(io.BytesIO(content))

        for page_number, page in enumerate(reader.pages, start=1):
            text = clean_text(page.extract_text() or "")

            if text:
                sections.append(
                    {
                        "text": text,
                        "page_number": page_number,
                        "source_location": f"Page {page_number}",
                    }
                )

    elif extension == ".docx":
        document = Document(io.BytesIO(content))

        parts = [paragraph.text for paragraph in document.paragraphs]

        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))

        text = clean_text("\n".join(parts))

        if text:
            sections.append(
                {
                    "text": text,
                    "source_location": "DOCX document",
                }
            )

    elif extension == ".pptx":
        presentation = Presentation(io.BytesIO(content))

        for slide_number, slide in enumerate(
            presentation.slides,
            start=1,
        ):
            parts = []

            for shape in slide.shapes:
                if shape.has_text_frame:
                    parts.append(shape.text)

                if shape.has_table:
                    for row in shape.table.rows:
                        parts.append(
                            " | ".join(
                                cell.text for cell in row.cells
                            )
                        )

            text = clean_text("\n".join(parts))

            if text:
                sections.append(
                    {
                        "text": text,
                        "source_location": f"Slide {slide_number}",
                    }
                )

    else:
        text = clean_text(
            content.decode("utf-8-sig", errors="replace")
        )

        if text:
            sections.append(
                {
                    "text": text,
                    "source_location": "TXT document",
                }
            )

    if not sections:
        raise ValueError(
            "No extractable text found. Scanned PDFs may need OCR."
        )

    return sections


def extract_text(filename: str, content: bytes) -> str:
    """Return all extracted text as a single string."""
    sections = extract_pages(filename, content)
    return "\n\n".join(section["text"] for section in sections)


def split_into_chunks(text: str) -> list[str]:
    """Split text into overlapping character-based chunks."""
    chunks = []
    start = 0

    while start < len(text):
        end = min(start + CHUNK_SIZE, len(text))
        chunk = text[start:end].strip()

        if chunk:
            chunks.append(chunk)

        if end == len(text):
            break

        start = end - CHUNK_OVERLAP

    return chunks


def create_embedding(text: str) -> list[float]:
    """Create an embedding using the local Ollama service."""
    url = f"{settings.ollama_base_url.rstrip('/')}/api/embed"

    response = requests.post(
        url,
        json={
            "model": EMBED_MODEL,
            "input": text,
        },
        timeout=settings.request_timeout_seconds,
    )
    response.raise_for_status()

    data = response.json()
    embeddings = data.get("embeddings", [])

    if not embeddings or not embeddings[0]:
        raise RuntimeError("Ollama returned an empty embedding.")

    return embeddings[0]


def ingest_document(filename: str, content: bytes) -> dict:
    """Extract, chunk, embed, and persist a document."""
    filename = Path(filename).name
    sections = extract_pages(filename, content)

    file_hash = hashlib.sha256(content).hexdigest()
    document_id = file_hash

    # Inspect existing chunks to support safe re-indexing of
    # documents indexed by the older version of this service.
    existing = collection.get(
        where={"document_id": document_id},
        include=["metadatas"],
    )

    existing_ids = existing.get("ids", [])
    existing_metadatas = existing.get("metadatas") or []

    if existing_ids:
        has_locations = all(
            metadata and "source_location" in metadata
            for metadata in existing_metadatas
        )

        if has_locations:
            return {
                "filename": filename,
                "document_id": document_id,
                "chunks": len(existing_ids),
                "status": "already_exists",
            }

        # Old chunks lack source-location metadata.
        # Remove only this document's old chunks before re-indexing.
        collection.delete(where={"document_id": document_id})

    ids = []
    documents = []
    metadatas = []
    embeddings = []

    chunk_index = 0

    for section in sections:
        section_chunks = split_into_chunks(section["text"])

        for chunk in section_chunks:
            ids.append(f"{document_id}_{chunk_index}")
            documents.append(chunk)

            metadata = {
                "document_id": document_id,
                "filename": filename,
                "chunk_index": chunk_index,
                "source_location": section["source_location"],
            }

            if "page_number" in section:
                metadata["page_number"] = section["page_number"]

            metadatas.append(metadata)
            chunk_index += 1

    # Generate embeddings before inserting new chunks.
    for chunk in documents:
        embeddings.append(create_embedding(chunk))

    collection.add(
        ids=ids,
        documents=documents,
        metadatas=metadatas,
        embeddings=embeddings,
    )

    return {
        "filename": filename,
        "document_id": document_id,
        "chunks": len(documents),
        "status": "indexed",
    }


def search_documents(query: str, top_k: int = 4) -> list[dict]:
    """Retrieve semantically relevant chunks and source locations."""
    if not query or not query.strip():
        raise ValueError("Search query cannot be empty.")

    if collection.count() == 0:
        return []

    query_embedding = create_embedding(query)

    results = collection.query(
        query_embeddings=[query_embedding],
        n_results=min(max(top_k, 1), collection.count()),
        include=["documents", "metadatas", "distances"],
    )

    matches = []

    docs = results.get("documents", [[]])[0] or []
    metadatas = results.get("metadatas", [[]])[0] or []
    distances = results.get("distances", [[]])[0] or []

    for doc, metadata, distance in zip(
        docs,
        metadatas,
        distances,
    ):
        metadata = metadata or {}

        match = {
            "text": doc,
            "filename": metadata.get("filename", "Unknown document"),
            "chunk_index": metadata.get("chunk_index", -1),
            "source_location": metadata.get(
                "source_location",
                "Location unavailable; re-upload to index it",
            ),
            "distance": distance,
        }

        if "page_number" in metadata:
            match["page_number"] = metadata["page_number"]

        matches.append(match)

    return matches
